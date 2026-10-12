"""
Monthly Menu PDF Generator - Epistemo Vikas Leadership School.

Reads menu items from menu.xlsx (sheets week1..week5) and generates a
month-wise menu PDF that replicates format.pdf, in this same folder.

Usage:
    python menu_app.py          # launches the GUI

Layout spec (measured from format.pdf):
    - Page : 841.92 x 595.32 pt (A4 landscape)
    - One page per Monday-to-Saturday week overlapping the chosen month.
      The first week is the one containing the 1st (e.g. if the 1st is a
      Tuesday, the week starts on Monday the 30th of the previous month;
      those previous-month days stay blank, without any HOLIDAY mark).
    - Week mapping : week1, week2, ... in order, rotating back to week1
      if a 6th week occurs.
    - Holiday     : 'HOLIDAY' in the first breakfast row, rest blank.
"""

import calendar
import datetime
import io
import os
try:
    import tkinter as tk
    from tkinter import ttk, messagebox
except ImportError:  # headless servers / Streamlit Cloud: core still works
    tk = ttk = messagebox = None
from xml.sax.saxutils import escape

import openpyxl
from PIL import Image
from pypdf import PdfReader
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.utils import ImageReader
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfgen import canvas as canvas_mod
from reportlab.platypus import Paragraph, Table, TableStyle

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
MENU_XLSX = os.path.join(BASE_DIR, "menu.xlsx")
FORMAT_PDF = os.path.join(BASE_DIR, "format.pdf")

# --------------------------------------------------------------------------
# Layout constants (measured from format.pdf)
# --------------------------------------------------------------------------
PAGE_W, PAGE_H = 841.92, 595.32

COL_WIDTHS = [134.7, 134.6, 134.7, 134.6, 134.6, 134.8]
TABLE_X = 17.0
TABLE_TOP = 125.2
TABLE_BOTTOM = 514.1

ROW_HEIGHTS = [
    27.1,                       # 0 title
    32.0,                       # 1 day names + dates
    18.4,                       # 2 breakfast section
    17.5, 17.6, 17.6, 17.5,     # 3-6 breakfast rows
    18.3,                       # 7 snacks section
    29.0,                       # 8 snacks row
    18.4,                       # 9 lunch section
    17.6, 17.5, 17.5, 17.6, 17.6, 17.5, 17.5, 17.5, 17.5, 17.7,  # 10-19 lunch rows
]

C_TITLE = HexColor("#7DB85B")
C_DAYS = HexColor("#8FBF6B")
C_SECTION = HexColor("#FFE7A8")
C_WHITE = HexColor("#FFFFFF")
C_ALT = HexColor("#EAF4E2")
C_GRID = HexColor("#C4D8B4")
C_BLACK = HexColor("#000000")

FONT = "Helvetica-Bold"
TITLE_SIZE = 17.5
DAY_SIZE = 12.25
SECTION_SIZE = 12.0
DATA_SIZE = 10.5

# Background image geometry inside format.pdf (full-width band)
BG_X0, BG_TOP, BG_W, BG_H = 0.0, 60.84, 841.92, 473.64
PX_PER_PT_X = 1024 / 841.92
PX_PER_PT_Y = 577 / 473.64

# Month pill inside the background image (source pixels)
PILL_PX = (522, 20, 763, 59)
PILL_FILL_PX = (528, 25, 758, 54)
PILL_FILL_COLOR = (162, 200, 145)
PILL_TEXT_COLOR = HexColor("#8E1F28")
PILL_FONT_SIZE = 19.5
PILL_MAX_TEXT_W = 188.0

DAY_NAMES = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday"]
MONTH_NAMES = list(calendar.month_name)[1:]


# --------------------------------------------------------------------------
# Data helpers
# --------------------------------------------------------------------------
def load_menu_data(path=MENU_XLSX):
    """Read menu.xlsx -> {sheet: {'sections': [...], 'breakfast': [...],
    'snack': [...], 'lunch': [...]}}; each data block is rows x 6 cols."""
    wb = openpyxl.load_workbook(path, data_only=True)
    data = {}
    for name in wb.sheetnames:
        ws = wb[name]
        def val(r, c):
            v = ws.cell(row=r, column=c).value
            if v is None:
                return ""
            v = str(v).strip()
            # Stale 'HOLIDAY' markers in the template are ignored;
            # holidays come only from the GUI inputs.
            if v.upper() == "HOLIDAY":
                return ""
            return v
        # Excel rows are 1-based: 4 breakfast sec, 5-8 bfast, 9 snack sec,
        # 10 snack, 11 lunch sec, 12-21 lunch.
        data[name.lower()] = {
            "sections": [val(4, 1), val(9, 1), val(11, 1)],
            "breakfast": [[val(r, c) for c in range(1, 7)] for r in range(5, 9)],
            "snack": [[val(10, c) for c in range(1, 7)]],
            "lunch": [[val(r, c) for c in range(1, 7)] for r in range(12, 22)],
        }
    return data


def monday_weeks(year, month):
    """Every Monday-to-Saturday block overlapping (year, month).

    The first week is the one containing the 1st, so it may start in the
    previous month (e.g. 1st = Tuesday -> week starts Monday 30th).
    If the 1st is a Sunday it belongs to no Mon-Sat block of this month,
    so the first week starts on Monday the 2nd.
    """
    first = datetime.date(year, month, 1)
    last = datetime.date(year, month, calendar.monthrange(year, month)[1])
    monday = first - datetime.timedelta(days=first.weekday())
    if monday + datetime.timedelta(days=5) < first:
        monday += datetime.timedelta(days=7)  # 1st is a Sunday
    weeks = []
    while monday <= last:
        weeks.append([monday + datetime.timedelta(days=i) for i in range(6)])
        monday += datetime.timedelta(days=7)
    return weeks


def sheet_for_week(week_index):
    """Sequential week1..week5; a 6th week rotates back to week1."""
    return "week%d" % ((week_index % 5) + 1)


def second_saturday(year, month):
    sats = [datetime.date(year, month, d)
            for d in range(1, calendar.monthrange(year, month)[1] + 1)
            if datetime.date(year, month, d).weekday() == 5]
    return sats[1] if len(sats) >= 2 else None


def parse_holiday_days(text, year, month):
    """Accept day numbers like '12, 13, 19, 20' (also '12/10', '12-10')."""
    days = set()
    bad = []
    last_day = calendar.monthrange(year, month)[1]
    for tok in text.replace(";", ",").replace(" ", ",").split(","):
        tok = tok.strip()
        if not tok:
            continue
        day = None
        if tok.isdigit():
            day = int(tok)
        else:
            for sep in ("/", "-", "."):
                if sep in tok:
                    part = tok.split(sep)[0].strip()
                    if part.isdigit():
                        day = int(part)
                    break
        if day is None or not 1 <= day <= last_day:
            bad.append(tok)
        else:
            days.add(day)
    return days, bad


def week_title(year, month, week_dates):
    start, end = week_dates[0], week_dates[-1]
    mon = calendar.month_name[month].upper()
    if start.month == end.month:
        rng = "%d-%d %s" % (start.day, end.day, start.strftime("%b"))
    else:
        rng = "%d %s-%d %s" % (start.day, start.strftime("%b"),
                               end.day, end.strftime("%b"))
    return "%s MENU \u2022 %s" % (mon, rng)


# --------------------------------------------------------------------------
# PDF building
# --------------------------------------------------------------------------
_bg_cache = {}


def get_background():
    """Background artwork extracted from format.pdf with the month pill
    interior repainted so the new month name can be drawn on top."""
    if "bg" in _bg_cache:
        return _bg_cache["bg"]
    reader = PdfReader(FORMAT_PDF)
    img = list(reader.pages[0].images)[0]
    bg = Image.open(io.BytesIO(img.data)).convert("RGB")
    x0, y0, x1, y1 = PILL_FILL_PX
    fill = Image.new("RGB", (x1 - x0, y1 - y0), PILL_FILL_COLOR)
    bg.paste(fill, (x0, y0))
    _bg_cache["bg"] = bg
    return bg


def pill_text_size(text):
    from reportlab.pdfbase.pdfmetrics import stringWidth
    w = stringWidth(text, FONT, PILL_FONT_SIZE)
    if w > PILL_MAX_TEXT_W:
        return PILL_FONT_SIZE * PILL_MAX_TEXT_W / w
    return PILL_FONT_SIZE


def build_pdf(year, month, holiday_dates, out_path):
    menu = load_menu_data()
    weeks = monday_weeks(year, month)
    if not weeks:
        raise ValueError("No Monday found in %s %d" % (MONTH_NAMES[month - 1], year))

    bg = get_background()
    bg_reader = ImageReader(bg)
    first_day = datetime.date(year, month, 1)

    c = canvas_mod.Canvas(out_path, pagesize=(PAGE_W, PAGE_H))

    st_title = ParagraphStyle("t", fontName=FONT, fontSize=TITLE_SIZE,
                              leading=TITLE_SIZE + 2, alignment=TA_CENTER,
                              textColor=C_BLACK)
    st_day = ParagraphStyle("d", fontName=FONT, fontSize=DAY_SIZE,
                            leading=DAY_SIZE + 2.5, alignment=TA_CENTER,
                            textColor=C_BLACK)
    st_sec = ParagraphStyle("s", fontName=FONT, fontSize=SECTION_SIZE,
                            leading=SECTION_SIZE + 2, alignment=TA_CENTER,
                            textColor=C_BLACK)
    st_data = ParagraphStyle("b", fontName=FONT, fontSize=DATA_SIZE,
                             leading=DATA_SIZE + 1.5, alignment=TA_CENTER,
                             textColor=C_BLACK)

    for wi, week in enumerate(weeks):
        sheet = sheet_for_week(wi)
        m = menu.get(sheet)
        if m is None:
            raise ValueError("Sheet '%s' not found in menu.xlsx" % sheet)

        # background artwork
        c.drawImage(bg_reader, BG_X0, PAGE_H - BG_TOP - BG_H,
                    width=BG_W, height=BG_H)

        # month pill text, e.g. OCTOBER MENU
        pill_text = "%s MENU" % calendar.month_name[month].upper()
        size = pill_text_size(pill_text)
        c.setFont(FONT, size)
        c.setFillColor(PILL_TEXT_COLOR)
        pill_cx = (PILL_PX[0] + PILL_PX[2]) / 2 / PX_PER_PT_X
        pill_cy_top = BG_TOP + (PILL_PX[1] + PILL_PX[3]) / 2 / PX_PER_PT_Y
        c.drawCentredString(pill_cx,
                            PAGE_H - (pill_cy_top + size * 0.36), pill_text)

        # which columns are holidays? Days spilling from the previous
        # month (before the 1st) stay blank with no HOLIDAY mark.
        is_pad = [d < first_day for d in week]
        is_hol = [d in holiday_dates for d in week]

        def cell(text):
            return Paragraph(escape(text) if text else "", st_data)

        # day header cells: Monday<br/>5/10/2026
        day_cells = [Paragraph("%s<br/>%d/%d/%d" % (DAY_NAMES[i], d.day,
                                                   d.month, d.year), st_day)
                     for i, d in enumerate(week)]

        # data blocks with holidays applied
        brow = []
        for r in range(4):
            row = []
            for i in range(6):
                if is_pad[i]:
                    row.append(cell(""))
                elif is_hol[i]:
                    row.append(cell("HOLIDAY" if r == 0 else ""))
                else:
                    row.append(cell(m["breakfast"][r][i]))
            brow.append(row)
        srow = [[cell("" if (is_pad[i] or is_hol[i]) else m["snack"][0][i])
                 for i in range(6)]]
        lrow = []
        for r in range(10):
            lrow.append([cell("" if (is_pad[i] or is_hol[i])
                              else m["lunch"][r][i])
                         for i in range(6)])

        table_data = (
            [[Paragraph(escape(week_title(year, month, week)), st_title)]
             + [""] * 5]
            + [day_cells]
            + [[Paragraph(escape(m["sections"][0]), st_sec)] + [""] * 5]
            + brow
            + [[Paragraph(escape(m["sections"][1]), st_sec)] + [""] * 5]
            + srow
            + [[Paragraph(escape(m["sections"][2]), st_sec)] + [""] * 5]
            + lrow
        )

        style_cmds = [
            ("SPAN", (0, 0), (-1, 0)),
            ("SPAN", (0, 2), (-1, 2)),
            ("SPAN", (0, 7), (-1, 7)),
            ("SPAN", (0, 9), (-1, 9)),
            ("BACKGROUND", (0, 0), (-1, 0), C_TITLE),
            ("BACKGROUND", (0, 1), (-1, 1), C_DAYS),
            ("BACKGROUND", (0, 2), (-1, 2), C_SECTION),
            ("BACKGROUND", (0, 7), (-1, 7), C_SECTION),
            ("BACKGROUND", (0, 9), (-1, 9), C_SECTION),
            ("GRID", (0, 0), (-1, -1), 0.75, C_GRID),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 2),
            ("RIGHTPADDING", (0, 0), (-1, -1), 2),
            ("TOPPADDING", (0, 0), (-1, -1), 1),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
        ]
        # alternating white / light-green data rows, restarted per section
        for r in (3, 4, 5, 6, 8) + tuple(range(10, 20)):
            if (r - (3 if r <= 6 else (8 if r == 8 else 10))) % 2 == 1:
                style_cmds.append(("BACKGROUND", (0, r), (-1, r), C_ALT))
            else:
                style_cmds.append(("BACKGROUND", (0, r), (-1, r), C_WHITE))

        t = Table(table_data, colWidths=COL_WIDTHS,
                  rowHeights=ROW_HEIGHTS)
        t.setStyle(TableStyle(style_cmds))
        t.wrapOn(c, 0, 0)
        t.drawOn(c, TABLE_X, PAGE_H - TABLE_BOTTOM)
        c.showPage()

    c.save()
    return out_path


# --------------------------------------------------------------------------
# GUI
# --------------------------------------------------------------------------
class MenuApp(tk.Tk if tk is not None else object):
    def __init__(self):
        super().__init__()
        self.title("Monthly Menu PDF Generator")
        self.resizable(False, False)
        now = datetime.date.today()

        pad = {"padx": 10, "pady": 5}
        frm = ttk.Frame(self, padding=12)
        frm.grid(row=0, column=0, sticky="nsew")

        ttk.Label(frm, text="Month:").grid(row=0, column=0, sticky="w", **pad)
        self.month_cb = ttk.Combobox(frm, values=MONTH_NAMES,
                                     state="readonly", width=15)
        self.month_cb.set(MONTH_NAMES[now.month - 1])
        self.month_cb.grid(row=0, column=1, sticky="w", **pad)
        self.month_cb.bind("<<ComboboxSelected>>", lambda e: self.refresh_info())

        ttk.Label(frm, text="Year:").grid(row=1, column=0, sticky="w", **pad)
        years = [str(y) for y in range(now.year - 2, now.year + 10)]
        self.year_cb = ttk.Combobox(frm, values=years,
                                    state="readonly", width=15)
        self.year_cb.set(str(now.year))
        self.year_cb.grid(row=1, column=1, sticky="w", **pad)
        self.year_cb.bind("<<ComboboxSelected>>", lambda e: self.refresh_info())

        ttk.Label(frm, text="Holidays (dates):").grid(row=2, column=0,
                                                      sticky="nw", **pad)
        self.hol_entry = ttk.Entry(frm, width=32)
        self.hol_entry.grid(row=2, column=1, sticky="w", **pad)
        self.hol_entry.bind("<KeyRelease>", lambda e: self.refresh_info())
        ttk.Label(frm, text='Day numbers separated by commas, e.g. 12, 13, 19, 20',
                  foreground="gray").grid(row=3, column=1, sticky="w",
                                          padx=10, pady=(0, 5))
        ttk.Label(frm, text='Days before the 1st are left blank (no menu, no HOLIDAY).',
                  foreground="gray").grid(row=4, column=1, sticky="w",
                                          padx=10, pady=(0, 5))

        self.sat_var = tk.BooleanVar(value=False)
        sat_chk = ttk.Checkbutton(frm, text="Second Saturday is a holiday",
                                  variable=self.sat_var,
                                  command=self.refresh_info)
        sat_chk.grid(row=5, column=0, columnspan=2, sticky="w", **pad)

        ttk.Label(frm, text="Weeks / holidays preview:").grid(
            row=6, column=0, sticky="nw", **pad)
        self.info = tk.Text(frm, width=46, height=8, state="disabled",
                            relief="sunken", borderwidth=1)
        self.info.grid(row=6, column=1, sticky="w", **pad)

        self.gen_btn = ttk.Button(frm, text="Generate PDF",
                                  command=self.on_generate)
        self.gen_btn.grid(row=7, column=0, columnspan=2, pady=12)

        self.status = ttk.Label(frm, text="", foreground="green")
        self.status.grid(row=8, column=0, columnspan=2, **pad)

        self.refresh_info()

    # -- helpers ------------------------------------------------------
    def selection(self):
        month = MONTH_NAMES.index(self.month_cb.get()) + 1
        return month, int(self.year_cb.get())

    def current_holidays(self):
        month, year = self.selection()
        days, bad = parse_holiday_days(self.hol_entry.get(), year, month)
        dates = {datetime.date(year, month, d) for d in days}
        if self.sat_var.get():
            ss = second_saturday(year, month)
            if ss:
                dates.add(ss)
        return dates, bad

    def refresh_info(self):
        try:
            month, year = self.selection()
        except ValueError:
            return
        dates, bad = self.current_holidays()
        weeks = monday_weeks(year, month)
        lines = []
        for wi, w in enumerate(weeks):
            tag = sheet_for_week(wi)
            rng = "%d %s-%d %s" % (w[0].day, w[0].strftime("%b"),
                                   w[-1].day, w[-1].strftime("%b"))
            hols = ["%d*" % d.day for d in w if d in dates]
            pad = " (blank before the 1st)" if w[0].month != month else ""
            lines.append("Week %d (%s): %s%s%s" % (
                wi + 1, tag, rng, pad,
                "  holidays: " + ",".join(hols) if hols else ""))
        if bad:
            lines.append("Ignored (invalid): " + ", ".join(bad))
        self.info.configure(state="normal")
        self.info.delete("1.0", "end")
        self.info.insert("1.0", "\n".join(lines))
        self.info.configure(state="disabled")

    def on_generate(self):
        try:
            month, year = self.selection()
        except ValueError:
            messagebox.showerror("Error", "Please select a valid month and year.")
            return
        dates, bad = self.current_holidays()
        if bad:
            if not messagebox.askyesno(
                    "Invalid dates",
                    "These holiday entries look invalid: %s\nContinue anyway?"
                    % ", ".join(bad)):
                return
        out = os.path.join(BASE_DIR, "%s %d Menu.pdf"
                           % (MONTH_NAMES[month - 1], year))
        try:
            build_pdf(year, month, dates, out)
        except Exception as ex:  # noqa: BLE001 - show any error to the user
            messagebox.showerror("Error", "Could not generate PDF:\n%s" % ex)
            return
        self.status.config(text="Saved: %s" % os.path.basename(out))
        messagebox.showinfo("Done", "Menu PDF saved in the men folder:\n%s" % out)


if __name__ == "__main__":
    if tk is None:
        raise SystemExit("tkinter is not available - install python3-tk to use the GUI.")
    MenuApp().mainloop()
